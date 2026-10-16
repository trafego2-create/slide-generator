"""Clona slides existentes (o 'modelo real') e troca texto/imagem preservando 100% da formatacao original."""
import copy
import re
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

_BOLD_MARKER_RE = re.compile(r"\*\*(.+?)\*\*")


def _parse_bold_segments(line):
    """Converte uma linha com marcadores **negrito** (ver parser.py) numa
    lista de (texto, bold). Sem marcador, retorna a linha inteira como um
    unico segmento nao-negrito."""
    segments = []
    pos = 0
    for m in _BOLD_MARKER_RE.finditer(line):
        if m.start() > pos:
            segments.append((line[pos:m.start()], False))
        segments.append((m.group(1), True))
        pos = m.end()
    if pos < len(line):
        segments.append((line[pos:], False))
    return segments or [(line, False)]


def duplicate_slide(prs, slide_index):
    """Duplica um slide (por indice) dentro da MESMA apresentacao, incluindo imagens.
    Retorna o novo objeto Slide, ja inserido no final do deck."""
    source = prs.slides[slide_index]
    dest = prs.slides.add_slide(source.slide_layout)

    # remove placeholders que add_slide criou automaticamente a partir do layout
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    rid_map = {}  # rId do slide de origem -> rId novo no slide de destino

    def remap_images(new_el):
        for blip in new_el.findall(".//" + qn("a:blip")):
            old_rid = blip.get(qn("r:embed"))
            if not old_rid:
                continue
            if old_rid not in rid_map:
                image_part = source.part.related_part(old_rid)
                new_rid = dest.part.relate_to(image_part, RT.IMAGE)
                rid_map[old_rid] = new_rid
            blip.set(qn("r:embed"), rid_map[old_rid])

    # fundo proprio do slide (ex: Atualidades V2 tem a barra do topo como
    # imagem de fundo do slide, nao do layout) -- sem isso o clone sai sem ela
    source_bg = source._element.cSld.bg
    if source_bg is not None:
        new_bg = copy.deepcopy(source_bg)
        remap_images(new_bg)
        dest_cSld = dest._element.cSld
        if dest_cSld.bg is not None:
            dest_cSld.remove(dest_cSld.bg)
        dest_cSld.insert(0, new_bg)

    # mapa de cores proprio do slide (ex: Atualidades V2 inverte texto/fundo,
    # e e isso que deixa o nome da categoria branco na etiqueta)
    source_clr = source._element.find(qn("p:clrMapOvr"))
    dest_clr = dest._element.find(qn("p:clrMapOvr"))
    if source_clr is not None and dest_clr is not None:
        dest_clr.getparent().replace(dest_clr, copy.deepcopy(source_clr))

    for shape in source.shapes:
        new_el = copy.deepcopy(shape._element)
        remap_images(new_el)
        dest.shapes._spTree.append(new_el)

    return dest


def find_shape_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"shape '{name}' nao encontrado no slide")


def set_label_text(shape, text):
    """Troca so o texto do ultimo run do primeiro paragrafo -- usado em
    etiquetas cujo comeco e decorativo (ex: as bolinhas da categoria em
    Atualidades), que precisam ficar intactas. Se o texto novo for maior que
    o do modelo, alarga a etiqueta na mesma proporcao (ex: 'GEOPOLÍTICA  |
    Oriente Médio' nao cabia na largura pensada pra '... |  América')."""
    runs = shape.text_frame.paragraphs[0].runs
    old_len = sum(len(r.text) for r in runs)
    runs[-1].text = text
    new_len = sum(len(r.text) for r in runs)
    if new_len > old_len:
        shape.width = int(shape.width * new_len / old_len)


def set_text_preserving_format(shape, text_block):
    """Substitui o texto do shape mantendo fonte/tamanho/negrito/cor do primeiro run
    original E as propriedades do paragrafo (bullet, indentacao, espacamento) em
    TODAS as linhas geradas -- sem isso, so a primeira linha ficaria com o marcador
    de lista e as demais sairiam sem bullet/indent. Trechos marcados com
    **negrito** (ver parser.py) viram runs em negrito dentro da mesma linha,
    sem alterar fonte/tamanho/cor herdados do modelo.

    A formatacao do run (rPr) e copiada inteira, nao atributo por atributo:
    cor de tema (ex: 'bg2 75%' no Atualidades V2) nao tem RGB e se perdia,
    deixando o texto na cor padrao -- que no V2 e branco no fundo branco."""
    tf = shape.text_frame
    first_para = tf.paragraphs[0]
    template_run = first_para.runs[0] if first_para.runs else None

    template_pPr = first_para._p.find(qn("a:pPr"))
    template_pPr = copy.deepcopy(template_pPr) if template_pPr is not None else None

    template_rPr = template_run._r.find(qn("a:rPr")) if template_run is not None else None
    font_bold = template_run.font.bold if template_run else None

    tf.clear()
    lines = text_block.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if template_pPr is not None:
            existing_pPr = p._p.find(qn("a:pPr"))
            if existing_pPr is not None:
                p._p.remove(existing_pPr)
            p._p.insert(0, copy.deepcopy(template_pPr))
        for text, is_bold in _parse_bold_segments(line):
            if not text:
                continue
            run = p.add_run()
            run.text = text
            if template_rPr is not None:
                existing_rPr = run._r.find(qn("a:rPr"))
                if existing_rPr is not None:
                    run._r.remove(existing_rPr)
                run._r.insert(0, copy.deepcopy(template_rPr))
            run.font.bold = True if is_bold else font_bold


def replace_picture(shape, new_image_path):
    """Troca so os bytes da imagem, mantendo posicao/tamanho/recorte (crop) do quadro original."""
    slide_part = shape.part
    image_part, new_rid = slide_part.get_or_add_image_part(new_image_path)
    blip = shape._element.find(".//" + qn("a:blip"))
    blip.set(qn("r:embed"), new_rid)
