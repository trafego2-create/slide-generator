"""Poda os 3 .pptx de modelo pra manter SO os slides que o renderer realmente
usa (capa, abertura, slide-modelo do clone). Todos os outros slides originais
sao descartados do resultado final de qualquer forma (ver renderer.py), entao
mante-los no arquivo fonte so desperdicava espaco -- o Atualidades_modelo.pptx
tinha 109MB so por causa de fotos de exemplo em slides que nunca aparecem na
saida. Isso tambem elimina a necessidade de Git LFS (o EasyPanel nao estava
resolvendo os ponteiros LFS no build, entao os pptx chegavam la corrompidos).

Roda uma vez: python scripts/prune_templates.py
Depois disso, ajustar config.py pros novos indices (sempre viram 0,1,2 pros
3 modelos, ja que keeper_indices=[0,1] e content_template_index vem depois)."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.oxml.ns import qn

from config import TEMPLATES


def keep_only_slides(prs, keep_indices):
    xml_slides = prs.slides._sldIdLst
    all_elements = list(xml_slides)
    keep_set = set(keep_indices)
    for i, el in enumerate(all_elements):
        if i not in keep_set:
            rid = el.get(qn("r:id"))
            xml_slides.remove(el)
            prs.part.drop_rel(rid)


def main():
    for key, cfg in TEMPLATES.items():
        path = cfg["file"]
        before_mb = path.stat().st_size / (1024 * 1024)

        prs = Presentation(str(path))
        keep = sorted(set(cfg["keeper_indices"] + [cfg["content_template_index"]]))
        print(f"{key}: mantendo indices originais {keep} (de {len(prs.slides._sldIdLst)} slides)")

        keep_only_slides(prs, keep)
        prs.save(str(path))

        after_mb = path.stat().st_size / (1024 * 1024)
        new_total = len(Presentation(str(path)).slides._sldIdLst)
        print(f"{key}: {before_mb:.1f}MB -> {after_mb:.1f}MB | slides finais: {new_total} (novos indices 0..{new_total-1})")


if __name__ == "__main__":
    main()
