"""Testa renderer.build_presentation contra os 3 modelos com conteudo sintetico,
sem depender de servidor web nem de chave da OpenAI. Roda: python scripts/smoke_test.py"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from renderer import build_presentation
from config import TEMPLATES

FAKE_PHOTO = r"C:\Users\trafe\Downloads\image.png"

LONG_TEXT = "\n".join([
    "Gestao de Continuidade de Negocios (GCN) e o conjunto de praticas de gestao que visa identificar potenciais ameacas a organizacao e os impactos que essas ameacas, caso se concretizem, podem causar as operacoes de negocio.",
    "O objetivo central da GCN e fornecer uma estrutura para aumentar a resiliencia organizacional, com a capacidade de uma resposta eficaz que salvaguarde os interesses das partes interessadas, a reputacao, a marca e as atividades de criacao de valor da organizacao.",
    "Entre os pilares da GCN estao: Analise de Impacto no Negocio (BIA), Avaliacao de Riscos, Estrategias de Continuidade, Planos de Resposta a Incidentes, Planos de Recuperacao de Desastres e Testes e Exercicios periodicos para validar a eficacia de todo o plano montado pela organizacao.",
    "E fundamental que a alta direcao esteja engajada no processo de GCN, pois sem patrocinio executivo os planos tendem a ficar desatualizados e a organizacao corre o risco de descobrir falhas criticas somente durante um incidente real.",
    "A Analise de Impacto no Negocio (BIA) e a etapa que identifica e avalia os efeitos potenciais de uma interrupcao nos processos criticos da organizacao, estabelecendo prioridades de recuperacao com base no tempo maximo tolerado de indisponibilidade de cada processo.",
    "Ja a Avaliacao de Riscos busca identificar eventos que possam causar interrupcoes, analisando a probabilidade de ocorrencia e o impacto potencial, para entao definir medidas de mitigacao proporcionais ao nivel de risco identificado pela equipe responsavel.",
])

SCENARIOS = {
    "petr": [
        {"corpo": LONG_TEXT, "imagem": FAKE_PHOTO},
        {"corpo": "Um segundo item, mais curto, so pra confirmar que nao sofre split desnecessario.", "imagem": FAKE_PHOTO},
    ],
    "bb": [
        {"corpo": "Enunciado: qual das alternativas abaixo esta correta? A) primeira B) segunda C) terceira"},
        {"corpo": "Comentario: a alternativa correta e a B, pois " + LONG_TEXT},
    ],
    "atualidades": [
        {
            "manchete": "GOVERNO ANUNCIA NOVA MEDIDA ECONOMICA",
            "resumo": LONG_TEXT,
            "imagem_1": FAKE_PHOTO,
            "imagem_2": FAKE_PHOTO,
        },
    ],
}


def main():
    for model_key, items in SCENARIOS.items():
        cfg = TEMPLATES[model_key]
        print(f"\n=== {model_key} ({cfg['label']}) ===")
        out_path = build_presentation(model_key, items, job_id="smoketest")
        size_mb = out_path.stat().st_size / (1024 * 1024)
        print(f"gerado: {out_path} ({size_mb:.2f} MB)")

        prs = Presentation(str(out_path))
        total = len(prs.slides._sldIdLst)
        print(f"total de slides no resultado: {total} (esperado: {len(cfg['keeper_indices'])} keepers + N gerados)")

        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.strip().replace("\n", " / ")[:50])
            print(f"  slide {i+1}: {texts}")


if __name__ == "__main__":
    main()
